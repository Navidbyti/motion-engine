import copy
import json
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from motion_engine.cli import main
from motion_engine.data_import import DataImportError, import_dataset
from motion_engine.qa import qa_report
from motion_engine.validation import load_spec, validate


def test_csv_import_builds_exact_mappings_that_pass_chart_qa(tmp_path):
    fragment = import_dataset(ROOT / "examples/assets/weather.csv", ROOT / "examples",
                              "temperatures", "weather_csv")
    assert fragment["dataset"]["rows"][0] == {"city": "Cairo", "temperatureC": 31}
    assert fragment["dataset"]["transform"]["columnSources"]["temperatureC"]["location"] == "B2:B4"
    spec = copy.deepcopy(load_spec(ROOT / "examples/weather.motion.json"))
    spec["sources"] = [fragment["source"]]
    spec["datasets"] = [fragment["dataset"]]
    assert validate(spec) == []
    assert not any(issue["code"].startswith("data_") for issue in qa_report(spec, ROOT / "examples")["issues"])
    output = tmp_path / "fragment.json"
    assert main(["import-data", str(ROOT / "examples/assets/weather.csv"), "--project-root", str(ROOT / "examples"),
                 "--dataset-id", "temperatures", "--source-id", "weather_csv", "--output", str(output)]) == 0
    assert json.loads(output.read_text(encoding="utf-8")) == fragment


def test_xlsx_import_renames_headers_and_preserves_numeric_and_text_cells(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Asset Data"
    sheet.append(["Asset Value", "Code", "Approved"])
    sheet.append([12.5, "001", True])
    sheet.append([13.25, "002", False])
    path = tmp_path / "data.xlsx"
    workbook.save(path)
    fragment = import_dataset(path, tmp_path, "assets", sheet="Asset Data")
    assert [column["type"] for column in fragment["dataset"]["columns"]] == ["number", "string", "boolean"]
    assert fragment["dataset"]["rows"][0] == {"Asset_Value": 12.5, "Code": "001", "Approved": True}
    assert fragment["dataset"]["transform"]["columnSources"]["Asset_Value"]["location"] == "'Asset Data'!A2:A3"
    assert fragment["source"]["uri"] == "data.xlsx"


def test_explicit_table_range_ignores_formatted_empty_rows_and_preserves_absolute_cells(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table"
    sheet["B3"] = "City"
    sheet["C3"] = "Value"
    sheet["B4"], sheet["C4"] = "A", 10
    sheet["B5"], sheet["C5"] = "B", 20
    sheet["D1000"] = "footer"
    path = tmp_path / "table.xlsx"
    workbook.save(path)
    fragment = import_dataset(path, tmp_path, "table_data", sheet="Table", table_range="B3:C5")
    assert fragment["dataset"]["rows"] == [{"City": "A", "Value": 10}, {"City": "B", "Value": 20}]
    assert fragment["dataset"]["transform"]["columnSources"]["Value"]["location"] == "'Table'!C4:C5"
    assert fragment["dataset"]["sourceRefs"][0]["location"] == "'Table'!B3:C5"


def test_import_refuses_ambiguous_blanks_uncalculated_formulas_and_unsafe_paths(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Values"
    sheet.append(["name", "value"])
    sheet.append(["A", 1])
    sheet.append(["B", None])
    path = tmp_path / "data.xlsx"
    workbook.save(path)
    with pytest.raises(DataImportError, match="blank cells"):
        import_dataset(path, tmp_path, "values", sheet="Values")
    sheet["B3"] = "=B2*2"
    workbook.save(path)
    with pytest.raises(DataImportError, match="no cached value"):
        import_dataset(path, tmp_path, "values", sheet="Values")
    with pytest.raises(DataImportError, match="inside"):
        import_dataset(path, tmp_path / "other", "values", sheet="Values")
