"""Source extraction with stable evidence locations and immutable hashes.

This module does not interpret the source's instructions or claims. It records
content and provenance for later review and MotionSpec authoring.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import subprocess
import wave
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Callable


Parser = Callable[[Path, str, int], tuple[list[dict[str, Any]], list[dict[str, str]]]]
PARSERS: dict[str, Parser] = {}
MEDIA_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
    ".wav": "audio/wav", ".mp3": "audio/mpeg", ".m4a": "audio/mp4",
    ".mp4": "video/mp4", ".mov": "video/quicktime", ".mkv": "video/x-matroska",
}


def register_parser(suffix: str, parser: Parser) -> None:
    """Register a parser for a lowercase extension; extensions may override it."""
    if not suffix.startswith(".") or suffix != suffix.lower():
        raise ValueError("parser suffix must be lowercase and begin with '.'")
    PARSERS[suffix] = parser


def _record(source_id: str, location: str, kind: str, value: Any,
            method: str, **metadata: Any) -> dict[str, Any]:
    result: dict[str, Any] = {
        "sourceId": source_id,
        "location": location,
        "kind": kind,
        "value": value,
        "method": method,
        "confidence": 1.0,
    }
    if metadata:
        result["metadata"] = metadata
    return result


def _text(path: Path, source_id: str, limit: int):
    records = []
    with path.open("r", encoding="utf-8-sig") as stream:
        for number, line in enumerate(stream, 1):
            if line.strip():
                records.append(_record(source_id, f"line:{number}", "text", line.rstrip("\r\n"), "utf-8 text"))
                _check_limit(records, limit)
    return records, []


def _csv(path: Path, source_id: str, limit: int):
    records = []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        try:
            headers = next(reader)
        except StopIteration:
            return [], [{"code": "empty_csv", "message": "CSV has no header row"}]
        if not headers or len(headers) != len(set(headers)) or any(not h.strip() for h in headers):
            raise ValueError("CSV headers must be present and unique")
        for col, header in enumerate(headers, 1):
            records.append(_record(source_id, f"{_column_name(col)}1", "header", header, "CSV parser"))
            _check_limit(records, limit)
        for row_number, row in enumerate(reader, 2):
            if len(row) != len(headers):
                raise ValueError(f"CSV row {row_number} has {len(row)} cells; expected {len(headers)}")
            for col, (header, value) in enumerate(zip(headers, row), 1):
                records.append(_record(source_id, f"{_column_name(col)}{row_number}", "cell", value, "CSV parser", column=header))
                _check_limit(records, limit)
    return records, []


def _pdf(path: Path, source_id: str, limit: int):
    from pypdf import PdfReader
    records = []
    issues = []
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ValueError("encrypted PDF requires an unlocked source")
    for page_number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if not text.strip():
            issues.append({"code": "page_text_empty", "message": f"Page {page_number} has no extractable text; OCR or visual review may be needed"})
        for line_number, line in enumerate(text.splitlines(), 1):
            if line.strip():
                records.append(_record(source_id, f"page:{page_number}/line:{line_number}", "text", line, "pypdf text extraction"))
                _check_limit(records, limit)
    return records, issues


def _docx(path: Path, source_id: str, limit: int):
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    records = []
    block_number = 0
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            block_number += 1
            value = Paragraph(child, document).text
            if value.strip():
                records.append(_record(source_id, f"block:{block_number}/paragraph", "text", value, "python-docx"))
                _check_limit(records, limit)
        elif isinstance(child, CT_Tbl):
            block_number += 1
            table = Table(child, document)
            for row_number, row in enumerate(table.rows, 1):
                for col_number, cell in enumerate(row.cells, 1):
                    records.append(_record(source_id, f"block:{block_number}/table:R{row_number}C{col_number}", "cell", cell.text, "python-docx"))
                    _check_limit(records, limit)
    return records, []


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _xlsx(path: Path, source_id: str, limit: int):
    from openpyxl import load_workbook
    records = []
    issues = []
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for sheet in formulas:
            cached_sheet = cached[sheet.title]
            for row, cached_row in zip(sheet.iter_rows(), cached_sheet.iter_rows()):
                for cell, cached_cell in zip(row, cached_row):
                    if cell.value is None:
                        continue
                    location = f"'{sheet.title.replace(chr(39), chr(39) * 2)}'!{cell.coordinate}"
                    if cell.data_type == "f":
                        value = str(cell.value)
                        cached_value = _json_value(cached_cell.value)
                        records.append(_record(source_id, location, "formula", value, "openpyxl", cachedValue=cached_value))
                        if cached_value is None:
                            issues.append({"code": "formula_cache_missing", "message": f"{location} has no cached result; recalculate in a spreadsheet application"})
                    else:
                        records.append(_record(source_id, location, "cell", _json_value(cell.value), "openpyxl", dataType=cell.data_type))
                    _check_limit(records, limit)
    finally:
        formulas.close()
        cached.close()
    return records, issues


def _image(path: Path, source_id: str, limit: int):
    from PIL import Image

    records = []
    try:
        with Image.open(path) as image:
            metadata = {
                "format": image.format, "width": image.width, "height": image.height,
                "mode": image.mode, "frameCount": getattr(image, "n_frames", 1),
                "exifOrientation": image.getexif().get(274),
            }
            profile = image.info.get("icc_profile")
            if profile:
                metadata["iccProfileSha256"] = hashlib.sha256(profile).hexdigest()
    except OSError as exc:
        raise ValueError(f"image {path} could not be read: {exc}") from exc
    for field, value in metadata.items():
        if value is not None:
            records.append(_record(source_id, f"image:{field}", "media_metadata", value, "Pillow header probe"))
            _check_limit(records, limit)
    issues = []
    if metadata["width"] * metadata["height"] > 50_000_000:
        issues.append({"code": "image_large", "message": "Image exceeds 50 megapixels; renderer may reject it"})
    return records, issues


def _wav(path: Path, source_id: str, limit: int):
    records = []
    try:
        with wave.open(str(path), "rb") as audio:
            rate, count = audio.getframerate(), audio.getnframes()
            metadata = {
                "channels": audio.getnchannels(), "sampleRateHz": rate,
                "sampleFrames": count, "sampleWidthBytes": audio.getsampwidth(),
                "duration": {"numerator": count, "denominator": rate},
            }
    except (OSError, EOFError, wave.Error) as exc:
        raise ValueError(f"WAV {path} could not be read: {exc}") from exc
    for field, value in metadata.items():
        records.append(_record(source_id, f"audio:{field}", "media_metadata", value, "Python wave header probe"))
        _check_limit(records, limit)
    return records, []


def _ffprobe(path: Path, source_id: str, limit: int):
    executable = shutil.which("ffprobe")
    if not executable:
        return [], [{"code": "ffprobe_unavailable", "message": "FFprobe is required to inspect compressed audio and video metadata"}]
    try:
        result = subprocess.run(
            [executable, "-v", "error", "-show_streams", "-show_format", "-of", "json", str(path)],
            capture_output=True, text=True, timeout=30, check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        return [], [{"code": "media_probe_failed", "message": f"FFprobe failed: {exc}"}]
    if result.returncode:
        return [], [{"code": "media_probe_failed", "message": f"FFprobe could not inspect media: {result.stderr.strip()[:500]}"}]
    try:
        payload = json.loads(result.stdout)
    except ValueError:
        return [], [{"code": "media_probe_failed", "message": "FFprobe returned invalid JSON"}]
    records = []
    for index, stream in enumerate(payload.get("streams", [])):
        for field in ("codec_type", "codec_name", "width", "height", "pix_fmt", "avg_frame_rate",
                      "sample_rate", "channels", "duration", "start_time", "nb_frames"):
            if stream.get(field) is not None:
                records.append(_record(source_id, f"stream:{index}/{field}", "media_metadata", stream[field], "FFprobe"))
                _check_limit(records, limit)
    duration = payload.get("format", {}).get("duration")
    if duration is not None:
        records.append(_record(source_id, "format:duration", "media_metadata", duration, "FFprobe"))
        _check_limit(records, limit)
    issues = [] if records else [{"code": "media_streams_empty", "message": "FFprobe found no usable stream metadata"}]
    return records, issues


def _column_name(number: int) -> str:
    name = ""
    while number:
        number, digit = divmod(number - 1, 26)
        name = chr(65 + digit) + name
    return name


def _check_limit(records: list[dict[str, Any]], limit: int) -> None:
    if len(records) > limit:
        raise ValueError(f"evidence limit {limit} exceeded; raise the limit explicitly")


def ingest(path: str | Path, source_id: str | None = None, limit: int = 100_000) -> dict[str, Any]:
    source_path = Path(path)
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    suffix = source_path.suffix.lower()
    parser = PARSERS.get(suffix)
    if parser is None:
        raise ValueError(f"unsupported source extension {suffix!r}; registered: {', '.join(sorted(PARSERS))}")
    if limit < 1:
        raise ValueError("limit must be positive")
    source_id = source_id or re.sub(r"[^A-Za-z0-9_.-]+", "_", source_path.stem)
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_.-]*", source_id):
        source_id = "source_" + source_id.lstrip("_0123456789")
    digest = hashlib.sha256()
    with source_path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    sha256 = digest.hexdigest()
    evidence, issues = parser(source_path, source_id, limit)
    return {
        "source": {"id": source_id, "uri": str(source_path), "mediaType": MEDIA_TYPES.get(suffix, "application/octet-stream"), "sha256": sha256},
        "evidence": evidence,
        "issues": issues,
        "recordCount": len(evidence),
    }


for _suffix, _parser in {
    ".txt": _text, ".md": _text, ".csv": _csv, ".pdf": _pdf,
    ".docx": _docx, ".xlsx": _xlsx,
    ".png": _image, ".jpg": _image, ".jpeg": _image, ".webp": _image,
    ".wav": _wav, ".mp3": _ffprobe, ".m4a": _ffprobe,
    ".mp4": _ffprobe, ".mov": _ffprobe, ".mkv": _ffprobe,
}.items():
    register_parser(_suffix, _parser)
