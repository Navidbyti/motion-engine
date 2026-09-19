"""Deterministic CSV/XLSX to MotionSpec dataset fragments with cell provenance."""
from __future__ import annotations

import csv
import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from .ingest import _column_name
from .revisions import file_sha256


class DataImportError(ValueError):
    pass


def _identifier(value: str, fallback: str) -> str:
    name = re.sub(r"[^A-Za-z0-9_.-]+", "_", value.strip()).strip("_.-")
    if not name or not name[0].isalpha():
        name = fallback if not name else "field_" + name
    return name


def _source_rows(path: Path, sheet: str | None, max_rows: int,
                 region: tuple[int, int, int, int] | None) -> tuple[list[list[Any]], str | None]:
    if path.suffix.lower() == ".csv":
        if sheet:
            raise DataImportError("CSV has no sheets")
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = []
            for number, row in enumerate(csv.reader(stream), 1):
                if region and number < region[0]:
                    continue
                if region and number > region[1]:
                    break
                rows.append(row[region[2] - 1:region[3]] if region else row)
                if len(rows) > max_rows + 1:
                    raise DataImportError(f"CSV exceeds {max_rows} data rows")
        return rows, None
    if path.suffix.lower() != ".xlsx":
        raise DataImportError("dataset import supports CSV and XLSX")
    if not sheet:
        raise DataImportError("XLSX import requires an explicit --sheet name")
    from openpyxl import load_workbook

    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if sheet not in formulas.sheetnames:
            raise DataImportError(f"sheet {sheet!r} does not exist")
        formula_sheet, value_sheet = formulas[sheet], cached[sheet]
        if not region and formula_sheet.max_row is not None and formula_sheet.max_row > max_rows + 1:
            raise DataImportError(f"sheet exceeds {max_rows} data rows")
        rows = []
        bounds = {"min_row": region[0], "max_row": region[1], "min_col": region[2], "max_col": region[3]} if region else {}
        for formula_row, value_row in zip(formula_sheet.iter_rows(**bounds), value_sheet.iter_rows(**bounds)):
            if len(rows) > max_rows:
                raise DataImportError(f"sheet exceeds {max_rows} data rows")
            result = []
            for original, resolved in zip(formula_row, value_row):
                if original.data_type == "f" and resolved.value is None:
                    raise DataImportError(f"formula {sheet}!{original.coordinate} has no cached value")
                result.append(resolved.value)
            rows.append(result)
        return rows, sheet
    finally:
        formulas.close()
        cached.close()


def _type_of(values: list[Any]) -> str:
    if all(isinstance(value, bool) or isinstance(value, str) and value.lower() in ("true", "false") for value in values):
        return "boolean"
    if all(isinstance(value, date) and (not isinstance(value, datetime) or value.time().isoformat() == "00:00:00") for value in values):
        return "date"
    if all(isinstance(value, int) and not isinstance(value, bool) or isinstance(value, str) and re.fullmatch(r"-?(?:0|[1-9]\d*)", value) for value in values):
        return "integer"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) or isinstance(value, str) and re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?", value) for value in values):
        return "number"
    return "string"


def _convert(value: Any, declared_type: str, location: str) -> Any:
    try:
        if declared_type == "string":
            return value.isoformat() if isinstance(value, (date, datetime)) else str(value)
        if declared_type == "boolean":
            if isinstance(value, bool):
                return value
            if isinstance(value, str) and value.lower() in ("true", "false"):
                return value.lower() == "true"
        if declared_type == "date":
            if isinstance(value, datetime):
                if value.time().isoformat() != "00:00:00":
                    raise ValueError("timestamp has a nonzero time component")
                return value.date().isoformat()
            if isinstance(value, date):
                return value.isoformat()
            return date.fromisoformat(value).isoformat()
        if declared_type == "integer":
            if isinstance(value, bool) or not re.fullmatch(r"-?(?:0|[1-9]\d*)", str(value)):
                raise ValueError("not a canonical integer")
            return int(value)
        if declared_type == "number":
            if isinstance(value, bool) or not re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?", str(value)):
                raise ValueError("not a canonical decimal")
            digits = re.sub(r"[^0-9]", "", str(value)).lstrip("0")
            if len(digits) > 15:
                raise ValueError("more than 15 significant digits; MotionSpec JSON numbers would lose precision")
            number = float(value)
            if not math.isfinite(number):
                raise ValueError("non-finite number")
            return number
    except (TypeError, ValueError) as exc:
        raise DataImportError(f"{location}: cannot convert {value!r} to {declared_type}: {exc}") from exc
    raise DataImportError(f"{location}: unsupported or invalid type {declared_type!r}")


def import_dataset(path: str | Path, project_root: str | Path, dataset_id: str,
                   source_id: str | None = None, sheet: str | None = None,
                   types: dict[str, str] | None = None, max_rows: int = 10_000,
                   table_range: str | None = None) -> dict[str, Any]:
    source_path = Path(path).resolve()
    root = Path(project_root).resolve()
    if not source_path.is_file() or not source_path.is_relative_to(root):
        raise DataImportError("source must be a file inside the future MotionSpec directory")
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_.-]*", dataset_id):
        raise DataImportError("dataset ID must be a MotionSpec identifier")
    source_id = source_id or _identifier(source_path.stem, "source")
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_.-]*", source_id):
        raise DataImportError("source ID must be a MotionSpec identifier")
    if source_id == dataset_id:
        raise DataImportError("source and dataset IDs must be distinct")
    if max_rows < 1:
        raise DataImportError("max_rows must be positive")
    region = None
    if table_range:
        match = re.fullmatch(r"([A-Z]+)([1-9]\d*):([A-Z]+)([1-9]\d*)", table_range)
        if not match:
            raise DataImportError("table range must look like A1:D12")
        first_col, first_row, last_col, last_row = match.groups()
        def index(column: str) -> int:
            number = 0
            for letter in column:
                number = number * 26 + ord(letter) - 64
            return number
        region = (int(first_row), int(last_row), index(first_col), index(last_col))
        if region[0] >= region[1] or region[2] > region[3] or region[1] - region[0] > max_rows:
            raise DataImportError("table range needs a header, data rows, and no more than max_rows rows")
    rows, selected_sheet = _source_rows(source_path, sheet, max_rows, region)
    if len(rows) < 2 or len(rows) - 1 > max_rows:
        raise DataImportError("source must contain a header and 1 to max_rows data rows")
    first_row = region[0] if region else 1
    first_col = region[2] if region else 1
    headers = rows[0]
    if region and len(headers) != region[3] - region[2] + 1:
        raise DataImportError("table range extends beyond available columns")
    if not headers or any(value is None or not str(value).strip() for value in headers):
        raise DataImportError("every selected column needs a header")
    if any(len(row) != len(headers) for row in rows):
        raise DataImportError("source has ragged rows; select a rectangular table")
    if any(value is None or value == "" for row in rows[1:] for value in row):
        raise DataImportError("blank cells need an explicit data policy before import")
    names = []
    for index, header in enumerate(headers, 1):
        name = _identifier(str(header), f"column_{index}")
        if name in names:
            base = f"{name}_{index}"
            name = base
            suffix = 2
            while name in names:
                name = f"{base}_{suffix}"
                suffix += 1
        names.append(name)
    overrides = types or {}
    if set(overrides) - set(names):
        raise DataImportError(f"type override names unknown columns: {sorted(set(overrides) - set(names))}")
    columns = []
    typed_rows = [dict() for _ in rows[1:]]
    mapping = {}
    sheet_prefix = f"'{selected_sheet.replace(chr(39), chr(39) * 2)}'!" if selected_sheet else ""
    for index, name in enumerate(names):
        values = [row[index] for row in rows[1:]]
        declared_type = overrides.get(name, _type_of(values))
        if declared_type not in ("string", "integer", "number", "boolean", "date"):
            raise DataImportError(f"column {name}: unsupported type {declared_type!r}")
        columns.append({"name": name, "type": declared_type})
        column_letter = _column_name(first_col + index)
        mapping[name] = {"sourceId": source_id, "location": f"{sheet_prefix}{column_letter}{first_row + 1}:{column_letter}{first_row + len(rows) - 1}"}
        for row_index, value in enumerate(values, first_row + 1):
            typed_rows[row_index - first_row - 1][name] = _convert(value, declared_type, f"{sheet_prefix}{column_letter}{row_index}")
    location = f"{sheet_prefix}{_column_name(first_col)}{first_row}:{_column_name(first_col + len(headers) - 1)}{first_row + len(rows) - 1}"
    media_type = "text/csv" if selected_sheet is None else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return {
        "source": {"id": source_id, "uri": source_path.relative_to(root).as_posix(), "mediaType": media_type,
                   "sha256": file_sha256(source_path), "authority": ["data"]},
        "dataset": {"id": dataset_id, "columns": columns, "rows": typed_rows,
                    "sourceRefs": [{"sourceId": source_id, "location": location, "method": "deterministic dataset import"}],
                    "transform": {"columnSources": mapping}},
        "headerMapping": [{"sourceHeader": str(header), "field": name} for header, name in zip(headers, names)],
    }
